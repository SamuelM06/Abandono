from datetime import datetime, date
from typing import List, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO
from app.schemas import CallDetail
from app.services import (
    get_kpis_for_date, get_kpis_for_datetime_range, get_hourly_stats,
    get_calls_for_export, ABANDONED_RESULTCALL, QUEUE_TIMEOUT_RESULTCALL,
    NO_ATENDIDA_RESULTDESC_IN,
)


XUMA_BLUE = "120180"
XUMA_GREEN_DARK = "00CD93"
XUMA_GREEN_LIGHT = "5AE280"
XUMA_GRAY = "333333"
WHITE = "FFFFFF"
LIGHT_GRAY = "F2F2F2"
LIGHT_BLUE = "E8EAF6"
LIGHT_GREEN = "E8F5E9"
ROJO = "C62828"
NARANJA = "E65100"


def _tiles_kpi_sin_atencion(kpis: dict) -> List[tuple]:
    return [
        ("LLAMADAS SIN ATENCION (A DEVOLVER)", kpis["total_ingresadas"], ROJO),
        ("EN ABANDONO", kpis.get("total_abandono", 0), XUMA_BLUE),
        ("FUERA DE HORARIO (OUT OF TIME IN)", kpis.get("total_fuera_horario", 0), NARANJA),
        ("QUEUE TIME OUT", kpis.get("total_queue_timeout", 0), XUMA_GRAY),
    ]


def create_kpi_sheet(
    wb: Workbook,
    kpis: dict,
    fecha: date,
    skills,
    report_title: str = "REPORTE ABANDONO XUMA - CONTINGENCIA",
    kpi_tiles: Optional[List[tuple]] = None,
    no_atendida: bool = False,
):
    ws = wb.active
    ws.title = "Resumen KPIs"

    header_font = Font(name="Raleway", bold=True, size=14, color=WHITE)
    header_fill = PatternFill(start_color=XUMA_BLUE, end_color=XUMA_BLUE, fill_type="solid")
    title_font = Font(name="Raleway", bold=True, size=18, color=XUMA_BLUE)
    subtitle_font = Font(name="Raleway", size=11, color=XUMA_GRAY)
    kpi_label_font = Font(name="Raleway", size=11, color=XUMA_GRAY)
    normal_font = Font(name="Raleway", size=11)

    thin_border = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )

    skills_label = ", ".join(skills) if isinstance(skills, (list, tuple)) else str(skills)

    ws.merge_cells(f"A1:{get_column_letter(6)}1")
    ws["A1"] = report_title
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells(f"A2:{get_column_letter(6)}2")
    ws["A2"] = f"Fecha: {fecha.strftime('%d/%m/%Y')} | Skill: {skills_label} | Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws["A2"].font = subtitle_font
    ws["A2"].alignment = Alignment(horizontal="center")

    if kpi_tiles is None:
        kpi_tiles = [
            ("TOTAL LLAMADAS INGRESADAS", kpis["total_ingresadas"], XUMA_BLUE),
            ("TOTAL LLAMADAS ABANDONO", kpis["total_abandono"], ROJO),
            ("TOTAL LLAMADAS ATENDIDAS", kpis["total_atendidas"], XUMA_GREEN_DARK),
        ]

    tiles_row = 4
    for i, (label, value, color) in enumerate(kpi_tiles):
        col = i * 2 + 1
        ws.merge_cells(start_row=tiles_row, start_column=col, end_row=tiles_row, end_column=col + 1)
        cell = ws.cell(row=tiles_row, column=col, value=value)
        cell.font = Font(name="Raleway", bold=True, size=28, color=color)
        cell.alignment = Alignment(horizontal="center")

        ws.merge_cells(start_row=tiles_row + 1, start_column=col, end_row=tiles_row + 1, end_column=col + 1)
        label_cell = ws.cell(row=tiles_row + 1, column=col, value=label)
        label_cell.font = kpi_label_font
        label_cell.alignment = Alignment(horizontal="center")

        for r in range(tiles_row, tiles_row + 2):
            for c in range(col, col + 2):
                ws.cell(row=r, column=c).border = thin_border
                ws.cell(row=r, column=c).fill = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")

    row = 7
    extra = [
        ("TOTAL UNICOS (primer registro por numero)", kpis.get("total_unicos", 0), XUMA_BLUE),
        ("TOTAL DUPLICADOS (rellamadas mismo numero)", kpis.get("total_duplicados", 0), NARANJA),
    ]
    for i, (label, value, color) in enumerate(extra):
        col = i * 3 + 1
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 2)
        cell = ws.cell(row=row, column=col, value=f"{label}: {value}")
        cell.font = Font(name="Raleway", bold=True, size=11, color=color)
        cell.alignment = Alignment(horizontal="center")

    ws.merge_cells(f"A9:{get_column_letter(6)}9")
    ws["A9"] = "DETALLE SIN ATENCION POR HORA" if no_atendida else "DETALLE POR HORA"
    ws["A9"].font = Font(name="Raleway", bold=True, size=14, color=XUMA_BLUE)

    if no_atendida:
        headers = ["HORA", "SIN ATENCION", "EN ABANDONO", "FUERA DE HORARIO", "QUEUE TIME OUT", "% DEL TOTAL"]
    else:
        headers = ["HORA", "TOTAL INGRESADAS", "ABANDONO", "ATENDIDAS", "% ABANDONO", "% ATENDIDAS"]

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=10, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = thin_border

    hourly = get_hourly_stats(fecha, skills, no_atendida=no_atendida)
    total_dia = sum(h["total"] for h in hourly) or 1

    for row_idx, h in enumerate(hourly, 11):
        total = h["total"] or 1
        ws.cell(row=row_idx, column=1, value=h["hora"]).font = normal_font

        if no_atendida:
            ws.cell(row=row_idx, column=2, value=h["total"]).font = normal_font
            ws.cell(row=row_idx, column=3, value=h["abandono"]).font = Font(name="Raleway", size=11, color=ROJO)
            ws.cell(row=row_idx, column=4, value=h["fuera_horario"]).font = Font(name="Raleway", size=11, color=NARANJA)
            ws.cell(row=row_idx, column=5, value=h["queue_timeout"]).font = Font(name="Raleway", size=11, color=XUMA_GRAY)
            c6 = ws.cell(row=row_idx, column=6, value=h["total"] / total_dia)
        else:
            ws.cell(row=row_idx, column=2, value=h["total"]).font = normal_font
            ws.cell(row=row_idx, column=3, value=h["abandono"]).font = Font(name="Raleway", size=11, color=ROJO)
            ws.cell(row=row_idx, column=4, value=h["atendidas"]).font = Font(name="Raleway", size=11, color=XUMA_GREEN_DARK)
            c5 = ws.cell(row=row_idx, column=5, value=h["abandono"] / total)
            c5.font = normal_font
            c5.number_format = "0.0%"
            c5.alignment = Alignment(horizontal="center")
            c6 = ws.cell(row=row_idx, column=6, value=h["atendidas"] / total)

        if no_atendida:
            c6.font = normal_font
        c6.number_format = "0.0%"
        c6.alignment = Alignment(horizontal="center")

        for col_idx in range(1, 7):
            ws.cell(row=row_idx, column=col_idx).border = thin_border
            if row_idx % 2 == 0:
                ws.cell(row=row_idx, column=col_idx).fill = PatternFill(start_color=LIGHT_BLUE, end_color=LIGHT_BLUE, fill_type="solid")

    totals_row = 11 + len(hourly)
    for col_idx in range(1, 7):
        ws.cell(row=totals_row, column=col_idx).font = Font(name="Raleway", bold=True, size=11, color=WHITE)
        ws.cell(row=totals_row, column=col_idx).fill = PatternFill(start_color=XUMA_GRAY, end_color=XUMA_GRAY, fill_type="solid")
        ws.cell(row=totals_row, column=col_idx).border = thin_border

    ws.cell(row=totals_row, column=1, value="TOTAL")
    ws.cell(row=totals_row, column=2).value = sum(h["total"] for h in hourly)
    ws.cell(row=totals_row, column=3).value = sum(h["abandono"] for h in hourly)

    if no_atendida:
        ws.cell(row=totals_row, column=4).value = sum(h["fuera_horario"] for h in hourly)
        ws.cell(row=totals_row, column=5).value = sum(h["queue_timeout"] for h in hourly)
    else:
        ws.cell(row=totals_row, column=4).value = sum(h["atendidas"] for h in hourly)

    for col_idx in range(1, 7):
        ws.column_dimensions[get_column_letter(col_idx)].width = 18


def create_detail_sheet(wb: Workbook, calls: List[CallDetail], fecha: date, time_range: str, title: str = "Detalle Llamadas"):
    ws = wb.create_sheet(title=title)

    header_font = Font(name="Raleway", bold=True, size=11, color=WHITE)
    header_fill = PatternFill(start_color=XUMA_BLUE, end_color=XUMA_BLUE, fill_type="solid")
    normal_font = Font(name="Raleway", size=10)
    thin_border = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )

    headers = [
        "ID", "FECHA", "HORA", "EXTENSIÓN", "AGENTE", "NOMBRE ASESOR", "SKILL",
        "TIPO LLAMADA", "CÓDIGO RESULTADO", "RESULTADO", "DETALLE RESULTADO",
        "DURACIÓN (seg)", "TIEMPO COLA (seg)", "NÚMERO LLAMANTE", "ANI LOCAL"
    ]

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = thin_border

    for row_idx, call in enumerate(calls, 2):
        ws.cell(row=row_idx, column=1, value=call.idlog_calls).font = normal_font
        ws.cell(row=row_idx, column=2, value=call.fecha.strftime("%d/%m/%Y")).font = normal_font
        ws.cell(row=row_idx, column=3, value=call.fecha.strftime("%H:%M:%S")).font = normal_font
        ws.cell(row=row_idx, column=4, value=call.extension or "").font = normal_font
        ws.cell(row=row_idx, column=5, value=call.agent or "").font = normal_font
        ws.cell(row=row_idx, column=6, value=call.asesor or "").font = normal_font
        ws.cell(row=row_idx, column=7, value=call.skill).font = normal_font
        ws.cell(row=row_idx, column=8, value=call.typecall or "").font = normal_font
        ws.cell(row=row_idx, column=9, value=call.resultcall or "").font = normal_font
        ws.cell(row=row_idx, column=10, value=call.resultcalldesc or "").font = normal_font
        ws.cell(row=row_idx, column=11, value=call.resultdesc or "").font = normal_font
        ws.cell(row=row_idx, column=12, value=round(call.timecall, 2) if call.timecall else 0).font = normal_font
        ws.cell(row=row_idx, column=13, value=round(call.timequeue, 2) if call.timequeue else 0).font = normal_font
        ws.cell(row=row_idx, column=14, value=call.numbercall or "").font = normal_font
        ws.cell(row=row_idx, column=15, value=call.localani or "").font = normal_font

        for col_idx in range(1, 16):
            ws.cell(row=row_idx, column=col_idx).border = thin_border
            if row_idx % 2 == 0:
                ws.cell(row=row_idx, column=col_idx).fill = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")

            if col_idx in [10, 11]:
                ws.cell(row=row_idx, column=col_idx).alignment = Alignment(wrap_text=True)

            if col_idx == 10 and call.resultcall == ABANDONED_RESULTCALL:
                ws.cell(row=row_idx, column=col_idx).font = Font(name="Raleway", size=10, bold=True, color=ROJO)
            elif col_idx == 10 and call.resultcall == "16":
                ws.cell(row=row_idx, column=col_idx).font = Font(name="Raleway", size=10, color=XUMA_GREEN_DARK)
            elif col_idx == 11 and (call.resultdesc or "").strip() in NO_ATENDIDA_RESULTDESC_IN:
                ws.cell(row=row_idx, column=col_idx).font = Font(name="Raleway", size=10, bold=True, color=NARANJA)
            elif col_idx == 10 and call.resultcall == QUEUE_TIMEOUT_RESULTCALL:
                ws.cell(row=row_idx, column=col_idx).font = Font(name="Raleway", size=10, color=XUMA_GRAY)

    column_widths = [8, 12, 10, 12, 14, 28, 22, 14, 14, 18, 30, 14, 14, 16, 16]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(calls) + 1}"
    ws.freeze_panes = "A2"


def generate_excel_report(
    kpis: dict,
    calls: List[CallDetail],
    fecha: date,
    skills,
    time_range: str,
    report_title: str = "REPORTE ABANDONO XUMA - CONTINGENCIA",
    detail_sheet_title: str = "Detalle Llamadas",
    no_atendida: bool = False,
) -> BytesIO:
    wb = Workbook()
    create_kpi_sheet(
        wb, kpis, fecha, skills, report_title,
        kpi_tiles=_tiles_kpi_sin_atencion(kpis) if no_atendida else None,
        no_atendida=no_atendida,
    )
    create_detail_sheet(wb, calls, fecha, time_range, detail_sheet_title)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_daily_excel_report(target_date: date, skills="In_Contingencias") -> BytesIO:
    kpis = get_kpis_for_date(target_date, skills)
    calls = get_calls_for_export(
        datetime.combine(target_date, datetime.min.time()),
        datetime.combine(target_date, datetime.max.time()),
        skills
    )
    return generate_excel_report(kpis, calls, target_date, skills, "Día completo")
